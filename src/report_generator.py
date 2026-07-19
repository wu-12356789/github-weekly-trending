from __future__ import annotations

from datetime import datetime
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from .models import Repository

BLUE=RGBColor(31,78,121); MUTED=RGBColor(90,105,120)
def set_font(run,name="Microsoft YaHei",size=10.5,bold=None,color=None):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),name); run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=color
def hyperlink(p,text,url):
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True); link=OxmlElement("w:hyperlink"); link.set(qn("r:id"),rid)
    r=OxmlElement("w:r"); rp=OxmlElement("w:rPr"); c=OxmlElement("w:color"); c.set(qn("w:val"),"0563C1"); u=OxmlElement("w:u"); u.set(qn("w:val"),"single"); rp.extend([c,u]); r.append(rp); t=OxmlElement("w:t"); t.text=text; r.append(t); link.append(r); p._p.append(link)
def page_field(p):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(); begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin"); instr=OxmlElement("w:instrText"); instr.text="PAGE"; sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate"); txt=OxmlElement("w:t"); txt.text="1"; end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end"); r._r.extend([begin,instr,sep,txt,end])
def cell_shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

class ReportGenerator:
    def __init__(self,output:Path): self.output=output; output.mkdir(parents=True,exist_ok=True)
    def generate(self,repos:list[Repository],when:datetime,images:dict[str,Path],history_available:bool):
        date=when.date().isoformat(); docx=self.output/f"GitHub每周热点项目推荐_{date}.docx"; md=self.output/f"GitHub每周热点项目推荐_{date}.md"
        self._docx(docx,repos,when,images,history_available); md.write_text(self._markdown(repos,when,history_available),encoding="utf-8"); return docx,md
    def _styles(self,doc):
        sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1); sec.header_distance=sec.footer_distance=Inches(.492)
        normal=doc.styles["Normal"]; normal.font.name="Microsoft YaHei"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); normal.font.size=Pt(10.5); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
        for name,size,before,after in [("Title",28,0,12),("Heading 1",16,16,8),("Heading 2",13,12,6),("Heading 3",11.5,8,4)]:
            s=doc.styles[name]; s.font.name="Microsoft YaHei"; s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei"); s.font.size=Pt(size); s.font.color.rgb=BLUE; s.font.bold=True; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
        page_field(sec.footer.paragraphs[0])
    def _docx(self,path,repos,when,images,history_available):
        doc=Document(); self._styles(doc); date=when.date().isoformat()
        if images.get("cover"): doc.add_picture(str(images["cover"]),width=Inches(6.5))
        p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("GitHub 每周热点项目推荐"),size=28,bold=True,color=BLUE)
        for text in [f"报告日期：{date}",f"统计范围：截至 {when:%Y-%m-%d %H:%M %Z} 的近 7 天", "由自动化系统生成；热度不代表安全性或生产可用性"]:
            p=doc.add_paragraph(text); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break(); doc.add_heading("目录",level=1)
        for text in ["1. 本周热点摘要","2. 热点项目总榜","3. 项目详细介绍","4. 本周趋势分析","5. 分类推荐","6. 风险提示","7. 数据来源和评分方法"]: doc.add_paragraph(text,style="List Number")
        doc.add_heading("1. 本周热点摘要",level=1)
        cats={r.category for r in repos}; top="、".join(r.name for r in repos[:3]); doc.add_paragraph(f"本周候选项目覆盖 {len(cats)} 个领域，热点集中在{'、'.join(list(cats)[:5])}。最值得优先关注的是 {top}。榜单兼顾近期增长、持续维护、README 完整度和实际使用价值，并通过分类配额避免单一领域占满。普通用户应优先选择安装路径清晰、带桌面界面或一键部署的项目；专业开发者可重点关注基础设施、Agent 与开发工具。{'本期周增长来自历史快照。' if history_available else '这是首次运行，周增长缺少基线，部分热度使用近期活跃度代理估算。'}")
        doc.add_heading("2. 热点项目总榜",level=1); table=doc.add_table(rows=1,cols=8); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
        widths=[.35,1.0,.9,.72,.65,.7,.62,1.56]
        heads=["排名","项目","分类","周增★","总★","语言","推荐","地址"]
        for i,(c,w,h) in enumerate(zip(table.rows[0].cells,widths,heads)): c.width=Inches(w); c.text=h; cell_shade(c,"D9EAF7"); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for i,r in enumerate(repos,1):
            cells=table.add_row().cells; vals=[str(i),r.name,r.category,str(r.weekly_stars),f"{r.stars:,}",r.language,str(r.analysis.get("ratings",{}).get("推荐指数",round(r.score/20,1))),""]
            for c,w,v in zip(cells,widths,vals): c.width=Inches(w); c.text=v; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            hyperlink(cells[-1].paragraphs[0],"GitHub",r.html_url)
        doc.add_heading("3. 项目详细介绍",level=1)
        for idx,r in enumerate(repos,1):
            if idx>1: doc.add_page_break()
            doc.add_heading(f"3.{idx} {r.name}",level=2); a=r.analysis
            if images.get(r.full_name): doc.add_picture(str(images[r.full_name]),width=Inches(6.3)); cap=doc.add_paragraph("非官方示意信息图"); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            meta=doc.add_table(rows=0,cols=2); meta.autofit=False
            pairs=[("仓库",r.full_name),("分类",r.category),("许可证 / 语言",f"{r.license} / {r.language}"),("Star / 周新增 / Fork",f"{r.stars:,} / {r.weekly_stars:,} / {r.forks:,}"),("最近更新 / Release",f"{r.updated_at.date()} / {r.latest_release}"),("平台 / 难度",f"{a.get('platforms','以官方说明为准')} / {a.get('difficulty','中等')}")]
            for label,value in pairs:
                cells=meta.add_row().cells; cells[0].width=Inches(1.55); cells[1].width=Inches(4.95); cells[0].text=label; cells[1].text=value; cell_shade(cells[0],"EDF3F8")
            sections=[("项目介绍",a.get("summary","")),("核心工作原理",a.get("principle","")),("为什么本周受到关注",a.get("why_hot","")),("适合人群",a.get("audience","")),("生产环境判断",a.get("production",""))]
            for h,t in sections: doc.add_heading(h,level=3); doc.add_paragraph(t)
            for h,key in [("核心功能","features"),("实际应用场景","scenarios"),("快速上手","quickstart"),("主要优点","pros"),("局限与风险","limitations")]:
                doc.add_heading(h,level=3)
                for x in a.get(key,[]): doc.add_paragraph(str(x),style="List Number" if key=="quickstart" else "List Bullet")
            doc.add_heading("评分与推荐理由",level=3); doc.add_paragraph("；".join(f"{k} {float(v):.1f}/5" for k,v in a.get("ratings",{}).items())+f"。综合热度算法得分 {r.score:.1f}/100。")
            p=doc.add_paragraph(); p.add_run("项目地址：").bold=True; hyperlink(p,r.html_url,r.html_url)
            p=doc.add_paragraph(); p.add_run("官方文档：").bold=True; hyperlink(p,r.docs_url or r.html_url,r.docs_url or r.html_url)
            if r.homepage: p=doc.add_paragraph(); p.add_run("官方网站：").bold=True; hyperlink(p,r.homepage,r.homepage)
        doc.add_page_break(); doc.add_heading("4. 本周趋势分析",level=1)
        languages={}; categories={}
        for r in repos: languages[r.language]=languages.get(r.language,0)+1; categories[r.category]=categories.get(r.category,0)+1
        doc.add_paragraph(f"领域分布：{categories}。语言分布：{languages}。AI Agent、本地模型与开发者效率工具仍是需要重点观察的方向；新项目依靠快速增长进入榜单，成熟项目则常因新版本重新升温。应结合周增长和 Release 时间判断，而不是只看总 Star。")
        doc.add_heading("5. 分类推荐",level=1)
        labels=["最适合学生","最适合普通用户","最适合程序员","最适合独立开发者","最适合企业团队","最值得长期关注"]
        for i,l in enumerate(labels): doc.add_paragraph(f"{l}：{repos[i%len(repos)].name} — {repos[i%len(repos)].analysis.get('summary','')}",style="List Bullet")
        doc.add_heading("6. 风险提示",level=1)
        for x in ["开源项目可能停止维护，热门不代表稳定。","安装第三方项目时核验发布者、依赖与脚本内容。","使用前检查许可证；企业使用前进行安全、合规和供应链审计。","不要直接运行来源不明的脚本，优先在隔离环境试用。"]: doc.add_paragraph(x,style="List Bullet")
        doc.add_heading("7. 数据来源和评分方法",level=1)
        doc.add_paragraph("数据来自 GitHub REST API 的 Search、README 与 Release 端点；采集时间见封面。综合分满分 100，权重覆盖周新增 Star、异常增速、总 Star、提交/Release/Issue/Fork 活跃、项目新鲜度、README、安装示例、维护与实用价值。GitHub Search 有索引延迟，Issue 数不等同活跃度，首次运行无历史基线，因此结果存在误差。")
        props=doc.core_properties; props.title="GitHub 每周热点项目推荐"; props.subject=date; doc.save(path)
    def _markdown(self,repos,when,history):
        lines=[f"# GitHub 每周热点项目推荐｜{when.date()}","",f"> 采集时间：{when.isoformat()}；{'已使用历史快照' if history else '首次运行，增长为估算/零基线'}。","","## 热点项目总榜","","|排名|项目|分类|周增 Star|总 Star|语言|得分|","|---:|---|---|---:|---:|---|---:|"]
        for i,r in enumerate(repos,1): lines.append(f"|{i}|[{r.name}]({r.html_url})|{r.category}|{r.weekly_stars}|{r.stars}|{r.language}|{r.score:.1f}|")
        for i,r in enumerate(repos,1):
            a=r.analysis; lines += ["",f"## {i}. {r.name}","",a.get("summary",""),"",f"- 分类：{r.category}",f"- GitHub：[{r.html_url}]({r.html_url})",f"- 官方文档：[{r.docs_url}]({r.docs_url})",f"- Star / 周增 / Fork：{r.stars} / {r.weekly_stars} / {r.forks}","","### 核心功能"]+[f"- {x}" for x in a.get("features",[])]+["","### 实用场景"]+[f"- {x}" for x in a.get("scenarios",[])]+["","### 优点与局限"]+[f"- 优点：{x}" for x in a.get("pros",[])]+[f"- 局限：{x}" for x in a.get("limitations",[])]
        lines += ["","## 风险提示","","热门不代表稳定或安全；安装前检查许可证、依赖和脚本，企业使用前进行安全审计。"]
        return "\n".join(lines)+"\n"
